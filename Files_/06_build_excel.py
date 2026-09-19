"""
Vendor Risk Scorecard workbook.
Tabs:
  1. Dashboard      - summary formulas (COUNTIFS/SUMIFS/AVERAGEIFS), no hardcoded results
  2. Vendor Scorecard - full scored population, conditional formatting by tier
  3. High Risk Detail - filtered High-tier vendors with recommended actions
  4. Data Dictionary  - column definitions + methodology note
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference

OUT = "/home/claude/vendor_risk_project/outputs/"
df = pd.read_csv(OUT + "vendors_scored.csv")

FONT = "Arial"
NAVY = "0C447C"
WHITE = "FFFFFF"
LOW_FILL = PatternFill("solid", fgColor="D9EAD3")
MED_FILL = PatternFill("solid", fgColor="FCE8B2")
HIGH_FILL = PatternFill("solid", fgColor="F4C7C3")
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
THIN = Side(style="thin", color="CCCCCC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

wb = Workbook()

# ============================================================
# TAB 1: DASHBOARD
# ============================================================
ws = wb.active
ws.title = "Dashboard"
ws.sheet_view.showGridLines = False

ws["B2"] = "Third-Party Vendor Risk Dashboard"
ws["B2"].font = Font(name=FONT, size=18, bold=True, color=NAVY)
ws["B3"] = "Source: Vendor Risk Tiering Model v1.0 \u2014 scored population on 'Vendor Scorecard' tab"
ws["B3"].font = Font(name=FONT, size=10, italic=True, color="666666")

headers = ["Metric", "Low", "Medium", "High", "Total"]
ws.append([])
start_row = 6
for j, h in enumerate(headers):
    c = ws.cell(row=start_row, column=2 + j, value=h)
    c.font = Font(name=FONT, bold=True, color=WHITE)
    c.fill = HEADER_FILL
    c.alignment = Alignment(horizontal="center")
    c.border = BORDER

n_vendors = len(df)
data_row_start = 2  # header row on Vendor Scorecard tab is row 1; data starts row 2
data_row_end = 1 + n_vendors

# Vendor count by tier
ws.cell(row=7, column=2, value="Vendor Count").font = Font(name=FONT, bold=True)
for j, tier in enumerate(["Low", "Medium", "High"]):
    col = get_column_letter(3 + j)
    ws.cell(row=7, column=3 + j,
             value=f'=COUNTIF(\'Vendor Scorecard\'!$N${data_row_start}:$N${data_row_end},"{tier}")')
ws.cell(row=7, column=6, value=f"=SUM(C7:E7)")

# Total annual spend by tier
ws.cell(row=8, column=2, value="Annual Spend ($)").font = Font(name=FONT, bold=True)
for j, tier in enumerate(["Low", "Medium", "High"]):
    ws.cell(row=8, column=3 + j,
             value=f'=SUMIFS(\'Vendor Scorecard\'!$F${data_row_start}:$F${data_row_end},'
                   f'\'Vendor Scorecard\'!$N${data_row_start}:$N${data_row_end},"{tier}")')
ws.cell(row=8, column=6, value="=SUM(C8:E8)")
for col in range(3, 7):
    ws.cell(row=8, column=col).number_format = '$#,##0'

# % of total spend by tier
ws.cell(row=9, column=2, value="% of Total Spend").font = Font(name=FONT, bold=True)
for j in range(3, 6):
    col = get_column_letter(j)
    ws.cell(row=9, column=j, value=f"={col}8/$F$8")
    ws.cell(row=9, column=j).number_format = "0.0%"
ws.cell(row=9, column=6, value="=F8/F8")
ws.cell(row=9, column=6).number_format = "0.0%"

# Average risk score by tier
ws.cell(row=10, column=2, value="Avg Risk Score (0-100)").font = Font(name=FONT, bold=True)
for j, tier in enumerate(["Low", "Medium", "High"]):
    ws.cell(row=10, column=3 + j,
             value=f'=AVERAGEIFS(\'Vendor Scorecard\'!$O${data_row_start}:$O${data_row_end},'
                   f'\'Vendor Scorecard\'!$N${data_row_start}:$N${data_row_end},"{tier}")')
    ws.cell(row=10, column=3 + j).number_format = "0.0"
ws.cell(row=10, column=6, value="=AVERAGE('Vendor Scorecard'!$O$2:$O$" + str(data_row_end) + ")")
ws.cell(row=10, column=6).number_format = "0.0"

# Vendors overdue for reassessment (>365 days) by tier
ws.cell(row=11, column=2, value="Overdue Reassessment (>365d)").font = Font(name=FONT, bold=True)
for j, tier in enumerate(["Low", "Medium", "High"]):
    ws.cell(row=11, column=3 + j,
             value=f'=COUNTIFS(\'Vendor Scorecard\'!$N${data_row_start}:$N${data_row_end},"{tier}",'
                   f'\'Vendor Scorecard\'!$T${data_row_start}:$T${data_row_end},">365")')
ws.cell(row=11, column=6, value="=SUM(C11:E11)")

for row in range(7, 12):
    for col in range(2, 7):
        ws.cell(row=row, column=col).border = BORDER
        if col > 2:
            ws.cell(row=row, column=col).alignment = Alignment(horizontal="center")

# Fill header tier colors as visual legend
ws.cell(row=7, column=3).fill = LOW_FILL
ws.cell(row=7, column=4).fill = MED_FILL
ws.cell(row=7, column=5).fill = HIGH_FILL

# Escalation flags
ws["B14"] = "Immediate Escalation Flags"
ws["B14"].font = Font(name=FONT, size=13, bold=True, color=NAVY)
ws["B15"] = "Sanctions screening flagged:"
ws["B15"].font = Font(name=FONT, bold=True)
ws["D15"] = f'=COUNTIF(\'Vendor Scorecard\'!$S${data_row_start}:$S${data_row_end},"Flagged")'
ws["B16"] = "Full data access + not certified:"
ws["B16"].font = Font(name=FONT, bold=True)
ws["D16"] = (f'=COUNTIFS(\'Vendor Scorecard\'!$I${data_row_start}:$I${data_row_end},"Full",'
             f'\'Vendor Scorecard\'!$K${data_row_start}:$K${data_row_end},"Not Certified")')
ws["B17"] = "2+ compliance incidents (3yr):"
ws["B17"].font = Font(name=FONT, bold=True)
ws["D17"] = f'=COUNTIF(\'Vendor Scorecard\'!$L${data_row_start}:$L${data_row_end},">=2")'

for r in range(15, 18):
    ws.cell(row=r, column=4).font = Font(name=FONT, bold=True, color="A32D2D")
    ws.cell(row=r, column=4).alignment = Alignment(horizontal="center")

# Bar chart: vendor count by tier
chart = BarChart()
chart.title = "Vendor Count by Risk Tier"
chart.y_axis.title = "Vendors"
chart.style = 10
data_ref = Reference(ws, min_col=3, max_col=5, min_row=7, max_row=7)
cats_ref = Reference(ws, min_col=3, max_col=5, min_row=6, max_row=6)
chart.add_data(data_ref, titles_from_data=False, from_rows=True)
chart.set_categories(cats_ref)
chart.series[0].tx = None
chart.legend = None
chart.varyColors = True
chart.height = 7
chart.width = 12
ws.add_chart(chart, "B20")

ws.column_dimensions["A"].width = 2
ws.column_dimensions["B"].width = 30
for col in "CDEF":
    ws.column_dimensions[col].width = 16

# ============================================================
# TAB 2: VENDOR SCORECARD (full population)
# ============================================================
ws2 = wb.create_sheet("Vendor Scorecard")
cols = ["vendor_id","vendor_name","industry","country","region_risk","annual_spend_usd",
        "contract_value_usd","relationship_years","data_access_level","business_criticality",
        "security_certification","compliance_incidents_3yr","sla_breach_count_1yr",
        "predicted_tier","risk_score_0_100","on_time_delivery_rate","subcontractor_use",
        "insurance_adequate","sanctions_screening","last_assessment_days_ago"]
header_map = {
    "vendor_id": "Vendor ID", "vendor_name": "Vendor Name", "industry": "Industry",
    "country": "Country", "region_risk": "Region Risk", "annual_spend_usd": "Annual Spend ($)",
    "contract_value_usd": "Contract Value ($)", "relationship_years": "Relationship (Yrs)",
    "data_access_level": "Data Access", "business_criticality": "Criticality",
    "security_certification": "Security Cert", "compliance_incidents_3yr": "Compliance Incidents (3yr)",
    "sla_breach_count_1yr": "SLA Breaches (1yr)", "predicted_tier": "Risk Tier",
    "risk_score_0_100": "Risk Score (0-100)", "on_time_delivery_rate": "On-Time Rate",
    "subcontractor_use": "Subcontractors", "insurance_adequate": "Insurance Adequate",
    "sanctions_screening": "Sanctions Screening", "last_assessment_days_ago": "Days Since Last Assessment",
}
out = df[cols].rename(columns=header_map)
for j, h in enumerate(out.columns, start=1):
    c = ws2.cell(row=1, column=j, value=h)
    c.font = Font(name=FONT, bold=True, color=WHITE, size=10)
    c.fill = HEADER_FILL
    c.alignment = Alignment(horizontal="center", wrap_text=True)
    c.border = BORDER
ws2.freeze_panes = "A2"

for i, row in enumerate(out.itertuples(index=False), start=2):
    for j, val in enumerate(row, start=1):
        c = ws2.cell(row=i, column=j, value=val)
        c.font = Font(name=FONT, size=10)
        c.border = BORDER
        if header_map["annual_spend_usd"] == out.columns[j-1] or header_map["contract_value_usd"] == out.columns[j-1]:
            c.number_format = '$#,##0'
        if out.columns[j-1] == header_map["risk_score_0_100"]:
            c.number_format = "0.0"
        if out.columns[j-1] == header_map["on_time_delivery_rate"]:
            c.number_format = "0.0%"

tier_col_idx = out.columns.get_loc(header_map["predicted_tier"]) + 1
tier_col_letter = get_column_letter(tier_col_idx)
rng_str = f"{tier_col_letter}2:{tier_col_letter}{1+n_vendors}"
ws2.conditional_formatting.add(rng_str, CellIsRule(operator="equal", formula=['"Low"'], fill=LOW_FILL))
ws2.conditional_formatting.add(rng_str, CellIsRule(operator="equal", formula=['"Medium"'], fill=MED_FILL))
ws2.conditional_formatting.add(rng_str, CellIsRule(operator="equal", formula=['"High"'], fill=HIGH_FILL))

widths = [11,20,16,14,11,14,15,12,11,11,13,13,12,10,13,10,13,13,15,14]
for j, w in enumerate(widths, start=1):
    ws2.column_dimensions[get_column_letter(j)].width = w

# ============================================================
# TAB 3: HIGH RISK DETAIL
# ============================================================
ws3 = wb.create_sheet("High Risk Detail")
hr = df[df["predicted_tier"] == "High"].sort_values("risk_score_0_100", ascending=False)
detail_cols = ["vendor_id","vendor_name","industry","predicted_tier","risk_score_0_100",
               "annual_spend_usd","data_access_level","security_certification",
               "sanctions_screening","compliance_incidents_3yr","last_assessment_days_ago"]
detail_header = {**header_map, "predicted_tier": "Risk Tier"}
hr_out = hr[detail_cols].rename(columns=detail_header)
hr_out["Recommended Action"] = hr_out.apply(
    lambda r: "Escalate to Compliance \u2014 sanctions flag" if r.get("Sanctions Screening") == "Flagged"
    else ("Remediate control gap \u2014 full data access, no certification"
          if r.get("Data Access") == "Full" and r.get("Security Cert") == "Not Certified"
          else ("Reassess \u2014 overdue >365 days" if r.get("Days Since Last Assessment") > 365
                else "Standard High-tier monitoring")),
    axis=1,
)

for j, h in enumerate(list(hr_out.columns), start=1):
    c = ws3.cell(row=1, column=j, value=h)
    c.font = Font(name=FONT, bold=True, color=WHITE, size=10)
    c.fill = PatternFill("solid", fgColor="A32D2D")
    c.alignment = Alignment(horizontal="center", wrap_text=True)
    c.border = BORDER
ws3.freeze_panes = "A2"
for i, row in enumerate(hr_out.itertuples(index=False), start=2):
    for j, val in enumerate(row, start=1):
        c = ws3.cell(row=i, column=j, value=val)
        c.font = Font(name=FONT, size=10)
        c.border = BORDER
widths3 = [11,20,16,10,12,14,11,13,15,13,14,32]
for j, w in enumerate(widths3, start=1):
    ws3.column_dimensions[get_column_letter(j)].width = w

# ============================================================
# TAB 4: DATA DICTIONARY
# ============================================================
ws4 = wb.create_sheet("Data Dictionary")
ws4["B2"] = "Data Dictionary & Methodology"
ws4["B2"].font = Font(name=FONT, size=16, bold=True, color=NAVY)

dict_rows = [
    ("Risk Tier", "Model-predicted tier (Low / Medium / High) from the Vendor Risk Tiering Model v1.0, a Random Forest classifier trained on 15 vendor attributes."),
    ("Risk Score (0-100)", "Blended score = 50 x P(Medium) + 100 x P(High), from the model's class probabilities. Higher = riskier."),
    ("Region Risk", "Geographic/regulatory risk tier of the vendor's home country (Low/Medium/High), based on standard country-risk categorization."),
    ("Data Access", "Level of company/customer data the vendor can access: No Access, Limited, or Full."),
    ("Security Cert", "Vendor's independent security certification status: Not Certified, SOC2, ISO27001, or Both."),
    ("Sanctions Screening", "Result of third-party sanctions/watchlist screening: Clear or Flagged."),
    ("Note", "All vendor data in this workbook is synthetically generated for a portfolio demonstration and does not represent real companies."),
]
r = 4
for label, desc in dict_rows:
    ws4.cell(row=r, column=2, value=label).font = Font(name=FONT, bold=True)
    ws4.cell(row=r, column=3, value=desc).font = Font(name=FONT)
    ws4.cell(row=r, column=3).alignment = Alignment(wrap_text=True, vertical="top")
    ws4.row_dimensions[r].height = 32
    r += 1
ws4.column_dimensions["B"].width = 20
ws4.column_dimensions["C"].width = 90

wb.save(OUT + "Vendor_Risk_Scorecard.xlsx")

# Print setup: landscape + fit to page width, so exports/prints render sensibly
for sheet in wb.worksheets:
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
wb.save(OUT + "Vendor_Risk_Scorecard.xlsx")
print("saved workbook")
